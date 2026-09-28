#!/usr/bin/env python3
"""반정형(MSTR 엠060003) 보고서 엑셀 다운로드 — 대체(새탭 PreferenceForm) 경로.

대상 예: 조달업체 면허 업종 등록 내역 (hubReptNo=00010, UI-ADOAAA-010R)
  - 팝업에서 `<input id="mf_popupCnts_btnExcelDown" type="button" value="엑셀다운로드">` 클릭
  - 새탭(MSTR 내보내기 옵션, PreferenceForm)이 열림
  - 새탭에서 `<input id="3131" type="submit" value="내보내기"
    onclick="disableButtonAndSubmit('3131','PreferenceForm');">` 클릭
  - 대기 후 .xlsx 다운로드

기존 정형(batch_extract_columns.py) 방식과의 차이:
  - 기존: scwin.btnExcelDown_onclick() JS 직접호출 → export 옵션창 expect_page
  - 신규: 엑셀다운로드 버튼 DOM 클릭 → 새탭 오픈 → submit 버튼 클릭 → 다운로드

사용법:
  uv run --with playwright --with openpyxl python tools/extract_semiform_columns.py --report-id 00010
  uv run --with playwright --with openpyxl python tools/extract_semiform_columns.py --report-id 00010 --headful
  uv run --with playwright --with openpyxl python tools/extract_semiform_columns.py --report-id 00010 --dry-run
"""
import argparse
import json
import sys
from pathlib import Path

BASE_DIR = Path(__file__).resolve().parent.parent
REPORTS_PATH = BASE_DIR / "public" / "reports131.json"
DATA_COLUMNS_DIR = BASE_DIR / "data" / "columns"
RAW_EXCEL_DIR = BASE_DIR / "data" / "raw_excel"

UA = ("Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
      "(KHTML, like Gecko) Chrome/126.0.0.0 Safari/537.36")

NAV_TIMEOUT = 60_000
SEARCH_TIMEOUT = 600_000
EXPORT_WINDOW_TIMEOUT = 60_000
DOWNLOAD_TIMEOUT = 300_000


def classify_column(col_name, sample_vals, is_numeric):
    name = col_name.strip()
    metric_keywords = ["금액", "수량", "단가", "건수", "율", "비율", "점수", "평균", "합계", "증감"]
    date_keywords = ["일자", "일시", "년월", "년도", "기간", "기한", "일"]
    code_keywords = ["코드", "번호", "차수", "사업자등록번호", "식별"]
    flag_keywords = ["여부", "구분", "유형", "상태"]
    if any(k in name for k in metric_keywords) or (is_numeric and not any(k in name for k in code_keywords + date_keywords)):
        return "metric"
    if any(k in name for k in code_keywords):
        return "identifier"
    if any(k in name for k in date_keywords):
        return "date"
    if any(k in name for k in flag_keywords):
        return "flag"
    return "dimension"


def parse_excel_file(excel_path, report_id, report_name, official_id):
    import openpyxl
    wb = openpyxl.load_workbook(str(excel_path), data_only=True)
    ws = wb.active
    header_row = None
    for r in range(1, min(15, ws.max_row + 1)):
        row_vals = [ws.cell(row=r, column=c).value for c in range(1, min(50, ws.max_column + 1))]
        non_empty = [v for v in row_vals if v is not None and str(v).strip()]
        if len(non_empty) >= 3 and not str(non_empty[0]).startswith("검색조건"):
            header_row = r
            break
    if not header_row:
        header_row = 5
    prompt_info = ws.cell(row=1, column=1).value or ""
    full_report_title = ws.cell(row=3, column=1).value or ""
    columns = []
    for c in range(1, ws.max_column + 1):
        col_name = ws.cell(row=header_row, column=c).value
        if col_name is None or not str(col_name).strip():
            continue
        col_name = str(col_name).strip()
        sample_vals = []
        for r in range(header_row + 1, min(header_row + 21, ws.max_row + 1)):
            v = ws.cell(row=r, column=c).value
            if v is not None:
                sample_vals.append(v)
        is_numeric = all(isinstance(v, (int, float)) for v in sample_vals) if sample_vals else False
        category = classify_column(col_name, sample_vals, is_numeric)
        columns.append({"index": c, "name": col_name, "category": category,
                        "is_metric": category == "metric", "is_numeric": is_numeric,
                        "sample_values": [str(x)[:40] for x in sample_vals[:3]]})
    metrics = [c["name"] for c in columns if c["is_metric"]]
    dimensions = [c["name"] for c in columns if c["category"] == "dimension"]
    identifiers = [c["name"] for c in columns if c["category"] == "identifier"]
    dates = [c["name"] for c in columns if c["category"] == "date"]
    return {"report_id": report_id, "report_name": report_name, "official_id": official_id,
            "full_report_title": full_report_title, "prompt_summary": prompt_info[:300] if prompt_info else "",
            "sheet_name": ws.title, "total_data_rows": max(0, ws.max_row - header_row),
            "total_columns": len(columns), "header_row_index": header_row,
            "metrics_summary": {"count": len(metrics), "metrics": metrics},
            "dimensions_summary": {"count": len(dimensions), "dimensions": dimensions},
            "identifiers_summary": {"count": len(identifiers), "identifiers": identifiers},
            "dates_summary": {"count": len(dates), "dates": dates},
            "columns": columns, "export_path": "semiform-newtab-PreferenceForm"}


def extract_semiform(report_id, headless=True, dry_run=False):
    from playwright.sync_api import sync_playwright
    reports = json.loads(REPORTS_PATH.read_text(encoding="utf-8"))
    report = next((r for r in reports if str(r.get("hubReptNo(reptId)")).strip() == report_id), None)
    if not report:
        raise ValueError(f"보고서 {report_id} 없음")
    report_name, official_id = report.get("보고서명"), report.get("보고서ID")
    print(f"[*] 대상: [{report_id}] {report_name} ({official_id})")
    DATA_COLUMNS_DIR.mkdir(parents=True, exist_ok=True)
    RAW_EXCEL_DIR.mkdir(parents=True, exist_ok=True)
    alerts = []
    with sync_playwright() as p:
        bw = p.chromium.launch(headless=headless, args=["--no-sandbox", "--disable-setuid-sandbox"])
        ctx = bw.new_context(user_agent=UA, viewport={"width": 1440, "height": 900}, accept_downloads=True)
        pg = ctx.new_page()
        pg.on("dialog", lambda d: (alerts.append(d.message), d.accept()))
        url = f"https://data.g2b.go.kr/link/AISC001_01/?reptNm={official_id}"
        print(f"[*] [1/5] 목록 접속: {url}")
        pg.goto(url, wait_until="networkidle", timeout=NAV_TIMEOUT)
        pg.wait_for_timeout(2000)
        title_sel = "#mf_wfm_container_reptLst_0_reptNm"
        try:
            pg.wait_for_selector(title_sel, timeout=10000)
        except Exception:
            title_sel = "a[id*='reptLst_0_reptNm']"
        print("[*] [2/5] 팝업 오픈")
        with ctx.expect_page(timeout=EXPORT_WINDOW_TIMEOUT) as pi:
            pg.click(title_sel)
        popup = pi.value
        popup.on("dialog", lambda d: (alerts.append(d.message), d.accept()))
        try:
            popup.wait_for_load_state("networkidle", timeout=NAV_TIMEOUT)
        except Exception:
            pass
        popup.wait_for_timeout(2000)
        print("[*] [3/5] 검색 실행")
        popup.wait_for_selector("#mf_popupCnts_btnS0001", timeout=15000)
        popup.click("#mf_popupCnts_btnS0001")
        popup.wait_for_timeout(2500)
        if alerts:
            print(f"[!] Alert: {' | '.join(alerts)}")
            bw.close()
            return {"status": "ALERT", "error": " | ".join(alerts)}
        print(f"[*] 조회완료 대기 (엑셀다운로드 .hide 제거, 최대 {SEARCH_TIMEOUT//1000}s)")
        try:
            popup.wait_for_selector("#mf_popupCnts_btnExcelDown:not(.hide)", timeout=SEARCH_TIMEOUT)
        except Exception:
            bw.close()
            return {"status": "TIMEOUT_SEARCH", "error": "엑셀 버튼 미활성화"}
        if dry_run:
            print("[+] Dry-Run OK")
            bw.close()
            return {"status": "DRY_RUN_OK"}
        # ── 핵심: 대체경로 — 엑셀다운로드 버튼 DOM 클릭 → 새탭 ──
        print("[*] [4/5] 엑셀다운로드 버튼 클릭 → 새탭(PreferenceForm) 대기")
        print(f"    버튼: input#mf_popupCnts_btnExcelDown.btn_cm.excel_down[value=엑셀다운로드]")
        with ctx.expect_page(timeout=EXPORT_WINDOW_TIMEOUT) as ei:
            # JS 우회 대신 실제 버튼 클릭 (새탭 window.open 유발)
            popup.click("#mf_popupCnts_btnExcelDown")
        export_page = ei.value
        export_page.on("dialog", lambda d: (alerts.append(d.message), d.accept()))
        try:
            export_page.wait_for_load_state("networkidle", timeout=NAV_TIMEOUT)
        except Exception:
            pass
        export_page.wait_for_timeout(2000)
        print(f"[*] 새탭 URL: {export_page.url}")
        print(f"[*] 새탭 타이틀: {export_page.title()}")
        # 새탭에서 내보내기 버튼 확인
        # 구문: <input onclick="disableButtonAndSubmit('3131','PreferenceForm');"
        #        name="3131" id="3131" type="submit" value="내보내기" class="mstrButton">
        try:
            export_page.wait_for_selector('input[id="3131"]', timeout=30000)
            print("[*] [5/5] 내보내기(submit#3131) 클릭 + 다운로드 대기")
        except Exception as e:
            html = export_page.content()[:2000]
            print(f"[!] 내보내기 버튼 없음. HTML 앞부분:\n{html}")
            bw.close()
            return {"status": "FAIL_NO_EXPORT_BTN", "error": str(e)[:200]}
        with export_page.expect_download(timeout=DOWNLOAD_TIMEOUT) as di:
            export_page.click('input[id="3131"]')
        dl = di.value
        excel_path = RAW_EXCEL_DIR / f"{report_id}_{dl.suggested_filename}"
        dl.save_as(str(excel_path))
        print(f"[+] 저장: {excel_path}")
        bw.close()
    import openpyxl  # noqa — 파싱 전 확인
    meta = parse_excel_file(excel_path, report_id, report_name, official_id)
    out = DATA_COLUMNS_DIR / f"{report_id}_columns.json"
    out.write_text(json.dumps(meta, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"[+] 메타 저장: {out} (컬럼 {meta['total_columns']}개, {meta['total_data_rows']}행)")
    return {"status": "SUCCESS", **{k: meta[k] for k in ("total_columns", "total_data_rows")}}


def main():
    ap = argparse.ArgumentParser(description="반정형 새탭(PreferenceForm) 엑셀 수집기")
    ap.add_argument("--report-id", default="00010")
    ap.add_argument("--headful", action="store_true")
    ap.add_argument("--dry-run", action="store_true")
    ap.add_argument("--force", action="store_true", help="기존 성공 산출물 무시하고 재실행")
    ap.add_argument("--all-semiform", action="store_true", help="mstrFrmt=엠060003 5종 일괄 (성공분 스킵)")
    a = ap.parse_args()
    if a.all_semiform:
        reports = json.loads(REPORTS_PATH.read_text(encoding="utf-8"))
        targets = [str(r.get("hubReptNo(reptId)")).strip() for r in reports if r.get("mstrFrmt") == "엠060003"]
        summary = {}
        for rid in targets:
            out = DATA_COLUMNS_DIR / f"{rid}_columns.json"
            if out.exists() and not a.force and not a.dry_run:
                print(f"[SKIP] {rid}: 기존 {out.name} 있어 스킵 (--force로 재실행 가능)")
                summary[rid] = "SKIPPED"
                continue
            res = extract_semiform(rid, headless=not a.headful, dry_run=a.dry_run)
            summary[rid] = res.get("status")
        print(json.dumps(summary, ensure_ascii=False, indent=2))
        return
    out = DATA_COLUMNS_DIR / f"{a.report_id}_columns.json"
    if out.exists() and not a.force and not a.dry_run:
        print(f"[SKIP] {a.report_id}: 기존 {out.name} 있어 스킵 (--force로 재실행 가능)")
        print(json.dumps({"status": "SKIPPED"}, ensure_ascii=False, indent=2))
        return
    res = extract_semiform(a.report_id, headless=not a.headful, dry_run=a.dry_run)
    print(json.dumps(res, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
