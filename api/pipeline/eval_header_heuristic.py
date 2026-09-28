#!/usr/bin/env python3
"""헤더/메타컬럼 휴리스틱 추출기 (읽기 전용 평가용 — 기존 로직 불변).

규칙:
  1. 첫 all-T 행 = 헤더행 (T=텍스트, 9=숫자, -=빈칸, 첫 8열 기준).
     단, 그 아래 행도 all-T면 타이틀/소계행이므로 스킵 (B형 대응).
  2. 헤더행 아래 데이터 행 스캔: 열별 텍스트 비율 높은 열 = 메타 컬럼.
  3. 메타 컬럼별 distinct 상위 N + 건수 출력.

출력: stdout 리포트만. 파일 쓰기 없음.
"""
import glob
import sys
from collections import Counter

SAMPLE_COLS = 8
SCAN_ROWS = 6
META_MIN_TEXT_RATIO = 0.5
TOP_N = 10
DATA_SCAN_LIMIT = 3000


def is_num(v):
    if v is None:
        return None
    s = str(v).strip().replace(",", "")
    if not s:
        return None
    try:
        float(s)
        return True
    except ValueError:
        return False


def row_marks(cells):
    out = []
    for c in cells:
        r = is_num(c)
        out.append("-" if r is None else ("9" if r else "T"))
    return "".join(out)


def main(top_n=TOP_N):
    from openpyxl import load_workbook

    files = sorted(glob.glob("data/raw_excel/*.xlsx"))
    for path in files:
        name = path.split("/")[-1]
        short = name[:24]
        try:
            wb = load_workbook(path, read_only=True, data_only=True)
        except Exception as e:
            print(f"== {short}: OPEN FAIL {type(e).__name__}")
            continue
        ws = wb.active
        try:
            head = []
            for i, row in enumerate(ws.iter_rows(values_only=True)):
                if i >= SCAN_ROWS:
                    break
                head.append(list(row)[:SAMPLE_COLS])
        except Exception as e:
            print(f"== {short}: HEAD PARSE ERR {str(e)[:50]}")
            wb.close()
            continue

        marks = [row_marks(r) for r in head]
        # 1. 헤더행 후보: 첫 all-T 행, 단 아래 행도 all-T면 스킵
        header_idx = None
        for i, m in enumerate(marks):
            if m and set(m) <= {"T"}:
                nxt = marks[i + 1] if i + 1 < len(marks) else ""
                if nxt and set(nxt) <= {"T"}:
                    continue  # 타이틀/소계행
                header_idx = i
                break

        print(f"== {short} (maxr={ws.max_row})")
        for i, m in enumerate(marks):
            tag = " <= HEADER" if i == header_idx else ""
            print(f"   row{i}: {m}{tag}")

        if header_idx is None:
            print("   -> 헤더 미검출 (C형: 데이터 없음 추정)")
            wb.close()
            continue

        headers = [(str(c).strip() if c is not None else "") for c in head[header_idx]]
        # 2. 데이터 행에서 열별 텍스트 비율
        text_counts: list[Counter] = [Counter() for _ in headers]
        n_data = 0
        try:
            for i, row in enumerate(ws.iter_rows(values_only=True)):
                if i <= header_idx:
                    continue
                cells = list(row)[: len(headers)]
                if all(c is None or str(c).strip() == "" for c in cells):
                    continue
                n_data += 1
                for j, c in enumerate(cells):
                    if j >= len(headers):
                        break
                    r = is_num(c)
                    if r is False:
                        text_counts[j][str(c).strip()] += 1
                if n_data >= DATA_SCAN_LIMIT:
                    break
        except Exception as e:
            print(f"   -> DATA SCAN ERR {str(e)[:60]}")
            wb.close()
            continue

        print(f"   -> data rows scanned: {n_data}")
        for j, h in enumerate(headers):
            if not h:
                continue
            total_text = sum(text_counts[j].values())
            ratio = total_text / max(n_data, 1)
            kind = "META" if ratio >= META_MIN_TEXT_RATIO else "num "
            top = text_counts[j].most_common(top_n)
            top_s = ", ".join(f"{v}({k})" for v, k in [(vv, kk) for kk, vv in top])
            # distinct 수 근사: Counter 크기 (스캔 범위 내)
            print(f"   [{kind} r={ratio:.2f} d~{len(text_counts[j])}] {h}: {top_s}")
        wb.close()


if __name__ == "__main__":
    n = int(sys.argv[1]) if len(sys.argv) > 1 else TOP_N
    main(top_n=n)
