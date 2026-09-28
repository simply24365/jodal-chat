"""xlsx 20개에서 메타 컬럼 distinct 수집 -> data/catalog/col_values.json.
규칙: 헤더행 아래 데이터에서 텍스트 비율>=0.5 열만, ID성(distinct/행수>0.8) 제외."""
import glob, json
from collections import Counter
from openpyxl import load_workbook

def is_num(v):
    if v is None: return None
    s = str(v).strip().replace(",", "")
    if not s: return None
    try: float(s); return True
    except: return False

# 00082/00065형: 첫 8열 컷에 걸리는 넓은 헤더 -> 24열로
SAMPLE_COLS = 24
SCAN_LIMIT = 5000
TOP_N = 30

out = {}
for path in sorted(glob.glob("data/raw_excel/*.xlsx")):
    rid = path.split("/")[-1][:5]
    try:
        wb = load_workbook(path, read_only=True, data_only=True)
    except Exception as e:
        print(f"{rid}: OPEN FAIL"); continue
    ws = wb.active
    head = []
    try:
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if i >= 8: break
            head.append(list(row)[:SAMPLE_COLS])
    except Exception:
        print(f"{rid}: HEAD ERR"); wb.close(); continue
    marks = []
    for ri, r in enumerate(head):
        m = ""
        for c in r:
            x = is_num(c)
            m += "-" if x is None else ("9" if x else "T")
        # 타이틀행: 앞 3열 중 2열 이상 빈칸이면 제목 병합행으로 스킵
        # (row0 검색조건, row2 보고서명 — 값은 첫 셀에만 있음)
        blanks_front = sum(1 for ch in m[:3] if ch == "-")
        if blanks_front >= 2:
            m = "TITLE"
        marks.append(m)
    header = None
    for i, m in enumerate(marks):
        if m == "TITLE":
            continue
        core = m.replace("-", "")
        if core and set(core) <= {"T"}:
            nxt = marks[i+1] if i+1 < len(marks) else ""
            nxtcore = nxt.replace("-", "")
            if nxtcore and set(nxtcore) <= {"T"}:
                continue
            header = i
            break
    if header is None:
        # 2행 헤더형: 연속 all-T 2행이면 두 번째를 헤더로 (00082/00065형)
        for i, m in enumerate(marks):
            if m == "TITLE":
                continue
            core = m.replace("-", "")
            nxt = marks[i+1] if i+1 < len(marks) else ""
            nxtcore = nxt.replace("-", "")
            if (core and set(core) <= {"T"} and nxtcore and set(nxtcore) <= {"T"}):
                header = i + 1
                break
    if header is None:
        print(f"{rid}: NO HEADER"); wb.close(); continue
    headers = [(str(c).strip() if c is not None else "") for c in head[header]]
    counters = [Counter() for _ in headers]
    n = 0
    try:
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if i <= header: continue
            cells = list(row)[:len(headers)]
            if all(c is None or str(c).strip() == "" for c in cells): continue
            n += 1
            for j, c in enumerate(cells):
                if j >= len(headers): break
                if is_num(c) is False:
                    counters[j][str(c).strip()] += 1
            if n >= SCAN_LIMIT: break
    except Exception as e:
        print(f"{rid}: SCAN ERR {str(e)[:50]}")
        wb.close(); continue
    wb.close()
    cols = {}
    for j, h in enumerate(headers):
        if not h: continue
        total = sum(counters[j].values())
        ratio = total / max(n, 1)
        if ratio < 0.5: continue
        distinct = len(counters[j])
        if distinct / max(n, 1) > 0.8: continue  # ID성
        top = counters[j].most_common(TOP_N)
        cols[h] = {"n": n, "distinct": distinct,
                   "top": [[v, k] for k, v in top]}
    out[rid] = {"header_row": header, "data_rows_scanned": n, "columns": cols}
    print(f"{rid}: header={header} rows={n} meta_cols={list(cols.keys())[:8]}")
json.dump(out, open("data/catalog/col_values.json", "w"), ensure_ascii=False, indent=1)
print("-> data/catalog/col_values.json", list(out.keys()))
